import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

EXCEL_DIR = os.path.join(os.path.expanduser("~"), "Downloads")
EXCEL_PATH = os.path.join(EXCEL_DIR, "incidents.xlsx")

HEADERS = [
    "Key",
    "User",
    "Reported Date by the Call Centre",
    "Reported Time by the Call Centre",
    "Customer",
    "Vehicle Stop Category",
    "Job No",
    "Vehicle No",
    "Driver Name",
    "Driver Contact No",
    "Vehicle Assigned by (Coordinator Name)",
    "Coordinator Mobile No",
    "Driver Contacted by OKI DOKI",
    "Driver Feedback - If Contacted",
    "Vehicle Parking with Goods",
    "Vehicle Stopped Location",
    "Vehicle Stopped Date",
    "Vehicle Stopped Time",
    "Vehicle Stopped Date & Time - V2",
    "Pickup Location",
    "Via Location/s",
    "Delivery Location",
    "Duration",
]


def _ensure_workbook():
    os.makedirs(EXCEL_DIR, exist_ok=True)
    if not os.path.exists(EXCEL_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Incidents"
        ws.append(HEADERS)
        wb.save(EXCEL_PATH)


def _combined_datetime(date_str, time_str):
    if not date_str or not time_str:
        return ""
    try:
        d = datetime.strptime(date_str, "%Y-%m-%d")
        return f"{d.strftime('%d-%m-%Y')} {time_str}:00"
    except ValueError:
        return f"{date_str} {time_str}"


def append_incident_row(incident, username: str, customer_name: str = ""):
    _ensure_workbook()

    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb["Incidents"]

        reported_at = incident.created_at or datetime.now()

        ws.append([
            incident.request_id,
            username,
            reported_at.strftime("%Y-%m-%d"),
            reported_at.strftime("%H:%M"),
            customer_name,
            incident.stop_category,
            incident.job_no,
            incident.vehicle_number,
            incident.driver_name,
            incident.driver_contact_number,
            incident.assigned_coordinator,
            incident.coordinator_mobile_number,
            incident.driver_contacted,
            incident.driver_feedback,
            "Yes" if incident.vehicle_parked else "No",
            incident.current_parking_location,
            incident.stopped_date,
            incident.stopped_time,
            _combined_datetime(incident.stopped_date, incident.stopped_time),
            incident.pickup_location,
            incident.via_locations,
            incident.delivery_location,
            incident.duration,
        ])

        wb.save(EXCEL_PATH)

    except PermissionError as exc:
        raise RuntimeError(
            "Could not write to incidents.xlsx — close the file in Excel and try again."
        ) from exc