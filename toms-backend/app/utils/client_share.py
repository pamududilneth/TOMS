import os
import re
from datetime import datetime
from openpyxl import Workbook, load_workbook

from .email_sender import send_report_email

CLIENT_REPORTS_DIR = os.path.join(
    os.path.expanduser("~"), "Downloads", "TOMS-client-reports"
)

HEADERS = [
    "Request ID",
    "Vehicle Number",
    "Driver Name",
    "Driver Contact Number",
    "Pickup Location",
    "Delivery Location",
    "Status",
    "Submitted At",
]


def _safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_\- ]", "", name).strip().replace(" ", "_")
    return cleaned or "client"


def _client_excel_path(client_name: str) -> str:
    os.makedirs(CLIENT_REPORTS_DIR, exist_ok=True)
    return os.path.join(CLIENT_REPORTS_DIR, f"{_safe_filename(client_name)}.xlsx")


def _ensure_client_workbook(path: str):
    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Incidents"
        ws.append(HEADERS)
        wb.save(path)


def _append_row(path: str, incident):
    _ensure_client_workbook(path)
    try:
        wb = load_workbook(path)
        ws = wb["Incidents"]
        ws.append([
            incident.request_id,
            incident.vehicle_number,
            incident.driver_name,
            incident.driver_contact_number,
            incident.pickup_location,
            incident.delivery_location,
            incident.status,
            incident.created_at.strftime("%Y-%m-%d %H:%M:%S")
            if incident.created_at else datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ])
        wb.save(path)
    except PermissionError as exc:
        raise RuntimeError(
            f"Could not write to {os.path.basename(path)} — close it in Excel and try again."
        ) from exc


def share_incident_with_client(client, incident, db):
    """
    Appends the incident to this client's personal Excel file every time.
    Emails the file only the first time this client is ever shared with,
    tracked via client.first_shared_at.
    """
    path = _client_excel_path(client.name)
    _append_row(path, incident)

    if client.first_shared_at is None and client.email:
        send_report_email(
            to_email=client.email,
            client_name=client.name,
            excel_path=path,
        )
        client.first_shared_at = datetime.now()
        db.add(client)
        db.commit()